"""Synthetic fixtures for txt/html/docx/tex import (no personal papers)."""
from pathlib import Path
import re
import pytest
from reader import store
from reader.importer import import_document


@pytest.fixture
def data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'DATA', tmp_path / 'data')
    return tmp_path


def test_txt_paragraphs(data_dir):
    p = data_dir / 'note.txt'
    p.write_text('First paragraph.\n\nSecond paragraph has more text.\n\n\nThird.\n', encoding='utf-8')
    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    assert texts == ['First paragraph.', 'Second paragraph has more text.', 'Third.']
    assert all(b['kind'] == 'paragraph' for b in d['blocks'])
    loc = d['atoms'][0]['location']
    assert loc['start'] == 0 and 'line_start' in loc and 'line_end' in loc
    assert d['source_file'] == 'source.txt'


def test_html_structure_strips_script_and_remote_image(data_dir):
    img = data_dir / 'local.png'
    img.write_bytes(
        b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde'
        b'\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82'
    )
    p = data_dir / 'article.html'
    p.write_text(
        '<html><head><style>.x{color:red}</style><script>alert(1)</script></head>'
        '<body><h1>Title One</h1><p>Hello body.</p>'
        '<img src="https://example.com/x.png" alt="remote">'
        '<img src="local.png" alt="local pic">'
        '<pre>code block</pre>'
        '<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>'
        '<ul><li>Item one</li></ul>'
        '</body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    kinds = [b['kind'] for b in d['blocks']]
    texts = [b['text'] for b in d['blocks']]
    assert 'Title One' in texts
    assert 'Hello body.' in texts
    assert any(b['kind'] == 'heading' for b in d['blocks'])
    assert any(b['kind'] == 'code' and 'code block' in b['text'] for b in d['blocks'])
    assert any(b['kind'] == 'table' and 'A' in b['text'] and '1' in b['text'] for b in d['blocks'])
    assert any(b['kind'] == 'paragraph' and 'Item one' in b['text'] for b in d['blocks'])
    assert not any('alert' in b['text'] for b in d['blocks'])
    assert any('remote' in (i['message'].lower()) or 'Image not embedded' in i['message'] for i in d['issues'])
    assert any(b.get('asset') for b in d['blocks'] if b['kind'] == 'figure')


def test_docx_heading_paragraph_table(data_dir):
    from docx import Document
    p = data_dir / 'paper.docx'
    doc = Document()
    doc.add_heading('Docx Title', level=1)
    doc.add_paragraph('Body paragraph about cooling.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).text = 'Value'
    table.cell(1, 0).text = 'PUE'
    table.cell(1, 1).text = '1.2'
    doc.save(p)
    d = import_document(p)
    assert any(b['kind'] == 'heading' and 'Docx Title' in b['text'] for b in d['blocks'])
    assert any(b['kind'] == 'paragraph' and 'cooling' in b['text'] for b in d['blocks'])
    assert any(b['kind'] == 'table' and 'PUE' in b['text'] for b in d['blocks'])
    loc = d['atoms'][0]['location']
    assert 'paragraph_index' in loc and loc['paragraph_index'] == 1
    assert 'start' in loc and 'end' in loc


def test_tex_section_input_comment_includegraphics(data_dir):
    p = data_dir / 'main.tex'
    p.write_text(
        r'''\documentclass{article}
% This comment should vanish
\section{Intro}
Plain paragraph with \textbf{bold} text.

\input{other.tex}
\includegraphics{fig.png}

\begin{equation}
E=mc^2
\end{equation}

\begin{verbatim}
raw code
\end{verbatim}
''',
        encoding='utf-8',
    )
    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    assert any(b['kind'] == 'heading' and b['text'] == 'Intro' for b in d['blocks'])
    assert any('bold' in b['text'] and 'comment' not in b['text'].lower() for b in d['blocks'])
    assert not any('This comment should vanish' in b['text'] for b in d['blocks'])
    assert any('not expanded' in i['message'] and 'other.tex' in i['message'] for i in d['issues'])
    assert any('includegraphics' in i['message'] for i in d['issues'])
    assert any(b['kind'] == 'formula' and 'E=mc^2' in b['text'].replace(' ', '') for b in d['blocks'])
    assert any(b['kind'] == 'code' and 'raw code' in b['text'] for b in d['blocks'])
    loc = d['atoms'][0]['location']
    assert 'line_start' in loc and 'start' in loc


def test_unsupported_extension(data_dir):
    p = data_dir / 'x.epub'
    p.write_bytes(b'PK fake')
    with pytest.raises(ValueError, match='Unsupported|Supported'):
        import_document(p)


def test_htm_alias(data_dir):
    p = data_dir / 'a.htm'
    p.write_text('<html><body><p>Hi there.</p></body></html>', encoding='utf-8')
    d = import_document(p)
    assert any('Hi there.' in b['text'] for b in d['blocks'])


def test_tex_full_document_wrapper(data_dir):
    """P1: begin{document}...end{document} must recurse, not collapse to one cleaned paragraph."""
    p = data_dir / 'wrapped.tex'
    content = (
        "\\documentclass{article}\n"
        "\\title{Measured Speeds}\n"
        "\\begin{document}\n"
        "\\section{Results}\n"
        "The measured speed is $v=\\frac{a}{b}$.\n"
        "\n"
        "\\begin{equation}\n"
        "E=mc^2\n"
        "\\end{equation}\n"
        "\n"
        "\\input{other.tex}\n"
        "\\includegraphics{fig.png}\n"
        "\\end{document}\n"
    )
    p.write_text(content, encoding='utf-8')
    d = import_document(p)
    # Title and/or section heading preserved (not wiped by document-env collapse).
    assert any(b['kind'] == 'heading' and ('Measured Speeds' in b['text'] or 'Results' in b['text']) for b in d['blocks'])
    assert any(b['kind'] == 'heading' and b['text'] == 'Results' for b in d['blocks'])
    # Fraction numerator/denominator not stripped by naive macro cleaning.
    para = next(b for b in d['blocks'] if 'The measured speed' in b['text'])
    assert '\\frac' in para['text']
    assert 'a' in para['text'] and 'b' in para['text']
    # Must not look like naive macro-strip residue (e.g. v= {t} / empty braces).
    assert '{t}' not in para['text']
    assert 'v= {}' not in para['text'].replace(' ', '')
    assert any(b['kind'] == 'formula' and 'E=mc^2' in b['text'].replace(' ', '') for b in d['blocks'])
    assert any('not expanded' in i['message'] and 'other.tex' in i['message'] for i in d['issues'])
    assert any('includegraphics' in i['message'] for i in d['issues'])
    # Must not collapse the whole document body into a single mangled paragraph.
    assert len(d['blocks']) >= 3


def test_html_div_article_span_body_text(data_dir):
    """P1: container tags with direct text must yield blocks."""
    p = data_dir / 'containers.html'
    p.write_text(
        '<html><body>'
        '<div>Visible body text.</div>'
        '<article>Article body.</article>'
        '<span>Span text.</span>'
        '<main>Main landmark text.</main>'
        '<section><p>Nested paragraph stays a block.</p>Trailing section text.</section>'
        '</body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    assert any('Visible body text.' in t for t in texts)
    assert any('Article body.' in t for t in texts)
    assert any('Span text.' in t for t in texts)
    assert any('Main landmark text.' in t for t in texts)
    assert any('Nested paragraph stays a block.' in t for t in texts)
    assert any('Trailing section text.' in t for t in texts)


def test_html_image_only_paragraph(data_dir):
    """P2: <p><img></p> must still produce a figure asset (not gated on paragraph text)."""
    from PIL import Image as PILImage
    from reader.store import folder
    img = data_dir / 'solo.png'
    PILImage.new('RGB', (4, 4), color=(0, 128, 255)).save(img)
    p = data_dir / 'img_only.html'
    p.write_text(
        '<html><body><p><img src="solo.png" alt="solo"></p></body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    figs = [b for b in d['blocks'] if b['kind'] == 'figure']
    assert len(figs) >= 1
    assert any(b.get('asset') for b in figs)
    assert any('solo' in b['text'] for b in figs)
    directory = folder(d['id'])
    assert any((directory / b['asset']).is_file() for b in figs if b.get('asset'))


def test_html_multi_img_paragraph(data_dir):
    """P2: multiple imgs in one <p> get distinct figure blocks/assets (no overwrite)."""
    from PIL import Image as PILImage
    from reader.store import folder
    a = data_dir / 'a.png'
    b = data_dir / 'b.png'
    PILImage.new('RGB', (4, 4), color=(255, 0, 0)).save(a)
    PILImage.new('RGB', (4, 4), color=(0, 255, 0)).save(b)
    p = data_dir / 'multi_img.html'
    p.write_text(
        '<html><body><p>Two pics<img src="a.png" alt="red"><img src="b.png" alt="green"></p></body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    assert any(b['kind'] == 'paragraph' and 'Two pics' in b['text'] for b in d['blocks'])
    figs = [b for b in d['blocks'] if b['kind'] == 'figure' and b.get('asset')]
    assert len(figs) >= 2
    assets = [b['asset'] for b in figs]
    assert len(set(assets)) == len(assets), assets
    directory = folder(d['id'])
    for name in assets:
        assert (directory / name).is_file()


def test_docx_heading_caption_math_picture(data_dir):
    """P2: drawings become figure assets; oMath is detected (text fallback and/or issue), never silent."""
    from docx import Document
    from docx.oxml import parse_xml
    from PIL import Image as PILImage
    from reader.store import folder

    img = data_dir / 'pic.png'
    PILImage.new('RGB', (8, 8), color=(200, 100, 50)).save(img)

    p = data_dir / 'rich.docx'
    doc = Document()
    doc.add_heading('Docx Math Pic Title', level=1)
    cap = doc.add_paragraph('Figure caption about the sample.')
    cap.add_run().add_picture(str(img))
    only = doc.add_paragraph()
    only.add_run().add_picture(str(img))
    omath = parse_xml(
        '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:oMath><m:r><m:t>E=mc2</m:t></m:r></m:oMath>'
        '</m:oMathPara>'
    )
    mp = doc.add_paragraph()
    mp._p.append(omath)
    doc.save(p)

    d = import_document(p)
    assert any(b['kind'] == 'heading' and 'Docx Math Pic Title' in b['text'] for b in d['blocks'])
    assert any('caption' in b['text'].lower() or 'Figure caption' in b['text'] for b in d['blocks'])
    figs = [b for b in d['blocks'] if b['kind'] == 'figure' and b.get('asset')]
    assert len(figs) >= 1
    directory = folder(d['id'])
    assert any((directory / b['asset']).is_file() for b in figs)
    math_blocks = [b for b in d['blocks'] if b['kind'] == 'formula' or 'E=mc2' in b['text'].replace(' ', '')]
    assert math_blocks, 'oMath must not be silently dropped'
    assert any('E=mc2' in b['text'].replace(' ', '') for b in math_blocks)


def test_html_comment_and_nonvisible_nodes_not_imported(data_dir):
    """P2 regression: Comment/Doctype/other non-visible strings are not body text."""
    p = data_dir / 'comment.html'
    p.write_text(
        '<html><body><!-- internal draft: do not publish --><p>Public article.</p></body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    assert [b['text'] for b in d['blocks']] == ['Public article.']

    other = data_dir / 'nonvisible.html'
    other.write_text(
        '<!DOCTYPE html><!-- top secret -->'
        '<html><body><![CDATA[hidden cdata]]><?pi secret?>'
        '<p>Visible only.</p></body></html>',
        encoding='utf-8',
    )
    d2 = import_document(other)
    assert [b['text'] for b in d2['blocks']] == ['Visible only.']


def test_html_inline_span_merges_into_one_paragraph(data_dir):
    p = data_dir / 'span_inline.html'
    p.write_text(
        '<html><body><div>The result is <span>not</span> significant.</div></body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    assert [b['text'] for b in d['blocks']] == ['The result is not significant.']


def test_tex_list_environment_preserves_math_and_issues(data_dir):
    """P1: itemize must recurse so \\frac survives and \\input/\\includegraphics emit issues."""
    p = data_dir / 'list.tex'
    p.write_text(
        "\\begin{itemize}\n"
        "\\item The speed is $v=\\frac{d}{t}$.\n"
        "\\input{missing}\n"
        "\\includegraphics{missing.png}\n"
        "\\end{itemize}\n"
        "\n"
        "\\begin{custombox}\n"
        "Unparsed $v=\\frac{a}{b}$ stays raw.\n"
        "\\end{custombox}\n",
        encoding='utf-8',
    )
    d = import_document(p)
    para = next(b for b in d['blocks'] if 'The speed is' in b['text'])
    assert r'\frac{d}{t}' in para['text']
    assert '{t}' not in para['text'].replace(r'\frac{d}{t}', '')
    assert any(
        'not expanded' in i['message'] and 'missing' in i['message'] for i in d['issues']
    )
    assert any(
        'includegraphics' in i['message'] and 'missing.png' in i['message'] for i in d['issues']
    )
    raw = next(b for b in d['blocks'] if r'\frac{a}{b}' in b['text'])
    assert 'stays raw' in raw['text']
    assert any(
        'custombox' in i['message'] and 'unresolved' in i['message'] and i['resolution'] is None
        for i in d['issues']
    )


def test_docx_table_cell_image_and_math(data_dir):
    """P2: table cells keep text plus drawing asset and oMath fallback (never silent)."""
    from docx import Document
    from docx.oxml import parse_xml
    from PIL import Image as PILImage
    from reader.store import folder

    img = data_dir / 'cell.png'
    PILImage.new('RGB', (6, 6), color=(1, 2, 3)).save(img)

    p = data_dir / 'cell.docx'
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = 'Cell content.'
    para = cell.paragraphs[0]
    para.add_run().add_picture(str(img))
    omath = parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>E=mc²</m:t></m:r></m:oMath>'
    )
    para._p.append(omath)
    doc.save(p)

    d = import_document(p)
    assert any(b['kind'] == 'table' and 'Cell content.' in b['text'] for b in d['blocks'])
    figs = [b for b in d['blocks'] if b['kind'] == 'figure' and b.get('asset')]
    assert figs, 'table cell image must become a figure asset'
    directory = folder(d['id'])
    assert any((directory / b['asset']).is_file() and (directory / b['asset']).stat().st_size > 0 for b in figs)
    table = next(b for b in d['blocks'] if b['kind'] == 'table')
    assert 'E=mc²' in table['text'].replace(' ', '') or 'E=mc^2' in table['text'].replace(' ', '')
    assert 'Cell content.' in table['text']


M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def _omath(xml):
    from docx.oxml import parse_xml
    return parse_xml(f'<m:oMath xmlns:m="{M_NS}">{xml}</m:oMath>')


def _omath_para(xml):
    from docx.oxml import parse_xml
    return parse_xml(
        f'<m:oMathPara xmlns:m="{M_NS}"><m:oMath>{xml}</m:oMath></m:oMathPara>'
    )


def test_docx_inline_math_keeps_sentence_order(data_dir):
    """P1: text runs and oMath stay in document order; fractions/scripts/grouping kept."""
    from docx import Document

    p = data_dir / 'inline_math.docx'
    doc = Document()
    para = doc.add_paragraph()
    para.add_run('Text A ')
    para._p.append(_omath('<m:sSub><m:e><m:r><m:t>h</m:t></m:r></m:e><m:sub><m:r><m:t>1</m:t></m:r></m:sub></m:sSub>'))
    para.add_run(' text B ')
    para._p.append(_omath('<m:sSub><m:e><m:r><m:t>p</m:t></m:r></m:e><m:sub><m:r><m:t>2</m:t></m:r></m:sub></m:sSub>'))
    para.add_run(' text C.')
    frac = doc.add_paragraph()
    frac.add_run('Ratio ')
    frac._p.append(_omath(
        '<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num>'
        '<m:den><m:r><m:t>b</m:t></m:r></m:den></m:f>'
    ))
    frac.add_run(' follows.')
    group = doc.add_paragraph()
    group.add_run('Group ')
    group._p.append(_omath(
        '<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
        '<m:e><m:r><m:t>x+y</m:t></m:r></m:e></m:d>'
    ))
    group.add_run('.')
    disp = doc.add_paragraph()
    disp._p.append(_omath_para(
        '<m:sSub><m:e><m:r><m:t>h</m:t></m:r></m:e><m:sub><m:r><m:t>0</m:t></m:r></m:sub></m:sSub>'
        '<m:r><m:t>=</m:t></m:r>'
        '<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>m</m:t></m:r></m:den></m:f>'
    ))
    table = doc.add_table(rows=1, cols=1)
    cell_para = table.cell(0, 0).paragraphs[0]
    cell_para.add_run('Cell ')
    cell_para._p.append(_omath(
        '<m:f><m:num><m:r><m:t>n</m:t></m:r></m:num>'
        '<m:den><m:r><m:t>m</m:t></m:r></m:den></m:f>'
    ))
    doc.save(p)

    d = import_document(p)
    inline = next(b for b in d['blocks'] if 'Text A' in b['text'])
    assert inline['kind'] == 'paragraph'
    assert re.search(r'Text A\s*\$.+\$\s*text B\s*\$.+\$\s*text C\.', inline['text'])
    assert 'h' in inline['text'] and 'p' in inline['text']
    assert not any(b['id'] != inline['id'] and b['kind'] == 'formula' and ('h_1' in b['text'] or b['text'] in ('h1', 'p2')) for b in d['blocks'])
    ratio = next(b for b in d['blocks'] if 'Ratio' in b['text'])
    assert r'\frac' in ratio['text'] and '{a}' in ratio['text'] and '{b}' in ratio['text']
    grouped = next(b for b in d['blocks'] if 'Group' in b['text'])
    assert '(x+y)' in grouped['text'].replace(' ', '') or r'(x+y)' in grouped['text']
    display = next(b for b in d['blocks'] if b['kind'] == 'formula')
    assert r'\frac' in display['text']
    cell = next(b for b in d['blocks'] if b['kind'] == 'table')
    assert 'Cell' in cell['text'] and r'\frac' in cell['text']


def test_docx_plain_super_subscript_runs(data_dir):
    """P1: vertAlign superscript/subscript stays distinct; mixed equals stays in place."""
    from docx import Document

    p = data_dir / 'scripts.docx'
    doc = Document()
    sq = doc.add_paragraph()
    sq.add_run('128')
    r = sq.add_run('2')
    r.font.superscript = True
    sq.add_run(' image.')
    cu = doc.add_paragraph()
    cu.add_run('128')
    r = cu.add_run('3')
    r.font.superscript = True
    cu.add_run(' volume.')
    mixed = doc.add_paragraph()
    mixed.add_run('38')
    r = mixed.add_run('2')
    r.font.superscript = True
    mixed._p.append(_omath('<m:r><m:t>=</m:t></m:r>'))
    mixed.add_run('1,444')
    sub = doc.add_paragraph()
    sub.add_run('H')
    r = sub.add_run('2')
    r.font.subscript = True
    sub.add_run('O.')
    doc.save(p)

    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    assert any('128²' in t for t in texts)
    assert any('128³' in t for t in texts)
    assert not any(re.search(r'1282\b', t) for t in texts)
    assert not any(re.search(r'1283\b', t) for t in texts)
    mixed_t = next(t for t in texts if '1,444' in t or '1444' in t)
    eq_at = mixed_t.find('$=$')
    if eq_at < 0:
        eq_at = mixed_t.find('=')
    assert eq_at > mixed_t.find('38')
    assert '1,444' in mixed_t[eq_at:]
    assert any('H₂O' in t or 'H$_{2}$O' in t or 'H$_2$O' in t for t in texts)


def test_tex_verb_not_treated_as_structure(data_dir):
    """P1: \\verb delimiters hide section/input/env tokens from the structure scan."""
    p = data_dir / 'verb.tex'
    p.write_text(
        'Tags look like: \\verb |\\section| or \\verb |\\paragraph|.\n'
        '\\section{Actual heading}\n'
        'Actual body.\n'
        '\n'
        'Also \\verb |\\input{secret.tex}| and \\verb |\\begin{document}| stay literal.\n',
        encoding='utf-8',
    )
    d = import_document(p)
    headings = [b for b in d['blocks'] if b['kind'] == 'heading']
    assert [b['text'] for b in headings] == ['Actual heading']
    sample = next(b for b in d['blocks'] if 'Tags look like' in b['text'])
    assert r'\section' in sample['text']
    assert r'\paragraph' in sample['text']
    assert sample['kind'] == 'paragraph'
    literal = next(b for b in d['blocks'] if 'stay literal' in b['text'])
    assert r'\input{secret.tex}' in literal['text']
    assert r'\begin{document}' in literal['text']
    assert not any('secret.tex' in i['message'] for i in d['issues'])
    assert not any(b['kind'] == 'heading' and 'paragraph' in b['text'].lower() for b in d['blocks'])


def test_tex_full_control_words_and_preamble_metadata(data_dir):
    """P2: titlerunning is not title; preamble macros stay out of body; metadata is kept."""
    wrapped = data_dir / 'meta.tex'
    wrapped.write_text(
        '\\documentclass{article}\n'
        '\\usepackage{graphicx}\n'
        '\\newcommand{\\foo}{bar}\n'
        '\\titlerunning{HTML papers on arXiv}\n'
        '\\title{Real Title}\n'
        '\\author{Ada Lovelace}\n'
        '\\begin{document}\n'
        '\\section{Body}\n'
        'Hello body.\n'
        '\\end{document}\n',
        encoding='utf-8',
    )
    d = import_document(wrapped)
    texts = [b['text'] for b in d['blocks']]
    headings = [b['text'] for b in d['blocks'] if b['kind'] == 'heading']
    assert 'Real Title' in headings
    assert 'Body' in headings
    assert not any('running{' in t or t.startswith('running') for t in texts)
    assert not any('usepackage' in t or 'newcommand' in t or '\\foo' in t for t in texts)
    assert any('Ada Lovelace' in t for t in texts)
    assert any('Hello body.' in t for t in texts)

    bare = data_dir / 'bare.tex'
    bare.write_text(
        '\\title{Outside Title}\n'
        '\\author{Grace Hopper}\n'
        '\\titlerunning{should not be title}\n'
        '\\newcommand{\\x}{y}\n'
        'Hello body without document env.\n',
        encoding='utf-8',
    )
    d2 = import_document(bare)
    headings2 = [b['text'] for b in d2['blocks'] if b['kind'] == 'heading']
    texts2 = [b['text'] for b in d2['blocks']]
    assert 'Outside Title' in headings2
    assert any('Grace Hopper' in t for t in texts2)
    assert not any('should not be title' in t for t in texts2)
    assert not any(t.startswith('running') for t in headings2)
    assert any('Hello body without document env.' in t for t in texts2)
    assert not any('\\x' == t or t.startswith('\\newcommand') for t in texts2)


def test_html_mathml_annotation_and_inline_spacing(data_dir):
    """P2: TeX logos stay unspaced; MathML annotation is not duplicated; keep real word spaces."""
    p = data_dir / 'mathml.html'
    p.write_text(
        '<html><body>'
        '<p>From <span class="ltx_TeX_logo" style="letter-spacing:-0.2em;">'
        'T<span>e</span>X</span> '
        '<math alttext="\\rightarrow" display="inline">'
        '<semantics><mo stretchy="false">→</mo>'
        '<annotation encoding="application/x-tex">\\rightarrow</annotation>'
        '</semantics></math> PDF.</p>'
        '<p>Keep <em>inline</em> space.</p>'
        '<p>Glue<span>d</span>word</p>'
        '</body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    logo = next(t for t in texts if 'PDF' in t)
    assert 'TeX' in logo
    assert 'T e X' not in logo
    assert logo.count('\\rightarrow') + logo.count('→') == 1
    assert not ('→' in logo and '\\rightarrow' in logo)
    assert any(t == 'Keep inline space.' for t in texts)
    assert any(t == 'Gluedword' for t in texts)


def test_docx_omath_para_keeps_every_sibling(data_dir):
    """P1: every m:oMath in an oMathPara is kept, in order, including alignment breaks."""
    from docx import Document
    from docx.oxml import parse_xml

    def para_of(parts):
        inner = ''.join(f'<m:oMath>{part}</m:oMath>' for part in parts)
        return parse_xml(f'<m:oMathPara xmlns:m="{M_NS}">{inner}</m:oMathPara>')

    def aligned(ch, num):
        return (
            f'<m:r><m:t>{ch}</m:t></m:r>'
            f'<m:r><m:rPr><m:aln/></m:rPr><m:t>={num}</m:t></m:r>'
        )

    p = data_dir / 'multi_omath.docx'
    doc = Document()
    doc.add_paragraph()._p.append(para_of([aligned('a', '1'), aligned('b', '2'), aligned('c', '3')]))
    doc.add_paragraph()._p.append(para_of([
        '<m:r><m:t>a=1</m:t></m:r>',
        '<m:r><m:t>b=2</m:t></m:r>',
    ]))
    doc.save(p)

    d = import_document(p)
    formulas = [b['text'] for b in d['blocks'] if b['kind'] == 'formula']
    assert len(formulas) >= 2
    aligned_line = next(t for t in formulas if 'c' in t and 'aligned' in t)
    compact = re.sub(r'\s+', '', aligned_line)
    assert r'\begin{aligned}' in aligned_line and r'\end{aligned}' in aligned_line
    assert 'a&=1' in compact and 'b&=2' in compact and 'c&=3' in compact
    assert compact.index('a&=1') < compact.index('b&=2') < compact.index('c&=3')
    plain = next(t for t in formulas if 'a=1' in re.sub(r'\s+', '', t))
    plain_c = re.sub(r'\s+', '', plain)
    assert 'b=2' in plain_c
    assert plain_c.index('a=1') < plain_c.index('b=2')
    assert r'\\' in plain


def test_docx_bar_pos_and_run_style(data_dir):
    """P2: m:bar pos top/bot; unsupported bar/style attributes become unresolved issues."""
    from docx import Document

    p = data_dir / 'bars.docx'
    doc = Document()
    doc.add_paragraph()._p.append(_omath(
        '<m:bar><m:e><m:r><m:t>x</m:t></m:r></m:e></m:bar>'
    ))
    doc.add_paragraph()._p.append(_omath(
        '<m:bar><m:barPr><m:pos m:val="top"/><m:ctrlPr/></m:barPr>'
        '<m:e><m:r><m:t>y</m:t></m:r></m:e></m:bar>'
    ))
    doc.add_paragraph()._p.append(_omath(
        '<m:bar><m:barPr><m:pos m:val="bot"/><m:ctrlPr/></m:barPr>'
        '<m:e><m:r><m:t>z</m:t></m:r></m:e></m:bar>'
    ))
    doc.add_paragraph()._p.append(_omath(
        '<m:bar><m:barPr><m:pos m:val="mid"/></m:barPr>'
        '<m:e><m:r><m:t>q</m:t></m:r></m:e></m:bar>'
    ))
    doc.add_paragraph()._p.append(_omath(
        '<m:bar><m:barPr><m:pos m:val="top"/><m:notARealProp m:val="1"/></m:barPr>'
        '<m:e><m:r><m:t>w</m:t></m:r></m:e></m:bar>'
    ))
    doc.add_paragraph()._p.append(_omath(
        '<m:r><m:rPr><m:sty m:val="b"/></m:rPr><m:t>v</m:t></m:r>'
    ))
    doc.add_paragraph()._p.append(_omath(
        '<m:r><m:rPr><m:sty m:val="bi"/></m:rPr><m:t>u</m:t></m:r>'
    ))
    doc.add_paragraph()._p.append(_omath(
        '<m:r><m:rPr><m:sty m:val="mystery"/></m:rPr><m:t>s</m:t></m:r>'
    ))
    doc.save(p)

    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    compact = [re.sub(r'\s+', '', t) for t in texts]
    assert any(r'\overline{x}' in t for t in compact)
    assert any(r'\overline{y}' in t for t in compact)
    assert any(r'\underline{z}' in t for t in compact)
    assert any(r'\mathbf{v}' in t for t in compact)
    assert any(r'\boldsymbol{u}' in t for t in compact)
    issues = [i for i in d['issues'] if i['id'].startswith('docx-omath')]
    assert len(issues) == 3
    assert all(i['resolution'] is None for i in issues)
    blob = '\n'.join(i['message'] for i in issues)
    assert 'mid' in blob
    assert 'notARealProp' in blob
    assert 'mystery' in blob
    assert 'm:val="bot"' not in blob
    assert any('q' in t for t in texts)
    assert any('w' in t for t in texts)


def test_html_block_boundaries_do_not_glue(data_dir):
    """P2: block edges are separated; inline spaces and TeX/LaTeX logo kerning stay as authored."""
    p = data_dir / 'blocks.html'
    p.write_text(
        '<html><body>'
        '<blockquote><p>First paragraph.</p><p>Second paragraph.</p></blockquote>'
        '<ul><li><p>Alpha.</p><p>Beta.</p></li></ul>'
        '<ul><li>Outer<ul><li>Inner item</li></ul></li></ul>'
        '<div><p>Left sentence.</p><p>Right sentence.</p></div>'
        '<p>Keep <em>inline</em> space.</p>'
        '<p>Glue<span>d</span>word</p>'
        '<p>From <span class="ltx_TeX_logo" style="letter-spacing:-0.2em;">'
        'T<span style="position:relative;bottom:-0.2ex;">e</span>X</span> and '
        '<span class="ltx_LaTeX_logo" style="letter-spacing:-0.2em;">'
        'L<span>a</span>T<span>e</span>X</span> stay tight.</p>'
        '</body></html>',
        encoding='utf-8',
    )
    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    blob = '\n'.join(texts)
    assert 'First paragraph.Second paragraph.' not in blob
    assert 'Alpha.Beta.' not in blob
    assert 'OuterInner' not in blob
    assert 'Left sentence.Right sentence.' not in blob
    assert 'First paragraph.' in blob and 'Second paragraph.' in blob
    assert 'Alpha.' in blob and 'Beta.' in blob
    assert 'Outer' in blob and 'Inner item' in blob
    assert any(t == 'Left sentence.' for t in texts)
    assert any(t == 'Right sentence.' for t in texts)
    assert any(t == 'Keep inline space.' for t in texts)
    assert any(t == 'Gluedword' for t in texts)
    logo = next(t for t in texts if 'stay tight' in t)
    assert 'TeX' in logo and 'T e X' not in logo
    assert 'LaTeX' in logo and 'L a T e X' not in logo


def test_tex_unclosed_brace_does_not_hang(data_dir):
    """P2: unclosed '{' / '[' in metadata arguments must finish and record an issue."""
    import signal
    from reader.importer import _tex_skip_groups

    src = r'\usepackage{foo'
    brace = src.index('{')

    def _alarm(signum, frame):
        raise TimeoutError('TeX brace skip hung on an unclosed group')

    old = signal.signal(signal.SIGALRM, _alarm)
    signal.alarm(5)
    try:
        after, ok = _tex_skip_groups(src, brace)
        assert ok is False
        assert after > brace
        closed = r'\newcommand{\foo}{bar} body'
        after_c, ok_c = _tex_skip_groups(closed, len(r'\newcommand'))
        assert ok_c is True
        assert closed[after_c:] == 'body'
        opted = r'\usepackage[utf8]{inputenc} NEXT'
        after_o, ok_o = _tex_skip_groups(opted, opted.index('['))
        assert ok_o is True
        assert opted[after_o:] == 'NEXT'

        p = data_dir / 'unclosed.tex'
        p.write_text(
            '\\documentclass{article}\n'
            '\\usepackage{foo\n'
            '\\title{Kept Title}\n'
            '\\author{Ada Partial\n'
            '\\hypersetup{colorlinks\n'
            '\\newcommand{\\zz\n'
            '\\includegraphics{fig.png\n'
            '\\input{missing.tex\n'
            '\\section[short\n'
            '\\section{After}\n'
            'Body survives the unclosed groups.\n',
            encoding='utf-8',
        )
        d = import_document(p)
    finally:
        signal.alarm(0)
        signal.signal(signal.SIGALRM, old)

    texts = [b['text'] for b in d['blocks']]
    headings = [b['text'] for b in d['blocks'] if b['kind'] == 'heading']
    assert 'Kept Title' in headings
    assert 'After' in headings
    assert any('Ada Partial' in t for t in texts)
    assert any('Body survives the unclosed groups.' in t for t in texts)
    assert not any('colorlinks' in t for t in texts)
    unclosed = [i for i in d['issues'] if 'unclosed' in i['message'].lower()]
    assert len(unclosed) >= 4
    assert all(i['resolution'] is None for i in unclosed)
    assert any('missing.tex' in i['message'] for i in d['issues'])
    assert any('includegraphics' in i['message'] for i in d['issues'])


def _sym_run(font, code):
    from docx.oxml import parse_xml
    return parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:sym w:font="{font}" w:char="{code}"/></w:r>'
    )


def test_docx_phantom_show_keeps_visible_base(data_dir):
    """P1: m:phant is visible unless m:show is 0/false/off. zeroAsc/Desc/Wid are not hide."""
    from docx import Document

    p = data_dir / 'phant.docx'
    doc = Document()
    # Real-like b00060: bar over (u–1), zeroAsc/zeroDesc, no m:show.
    doc.add_paragraph()._p.append(_omath(
        '<m:phant><m:phantPr><m:zeroAsc m:val="1"/><m:zeroDesc m:val="1"/></m:phantPr>'
        '<m:e><m:acc><m:accPr><m:chr m:val="̅"/></m:accPr>'
        '<m:e><m:r><m:t>u</m:t></m:r></m:e></m:acc>'
        '<m:r><m:rPr><m:nor/></m:rPr><m:t>–</m:t></m:r>'
        '<m:r><m:t>1</m:t></m:r></m:e></m:phant>'
    ))
    # Real-like b00110: zeroWid phantom still shows C(Φ[q]).
    doc.add_paragraph()._p.append(_omath(
        '<m:phant><m:phantPr><m:zeroWid m:val="1"/></m:phantPr><m:e>'
        '<m:r><m:t>C</m:t></m:r>'
        '<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>'
        '<m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>Φ</m:t></m:r>'
        '<m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr>'
        '<m:e><m:r><m:t>q</m:t></m:r></m:e></m:d>'
        '</m:e></m:d></m:e></m:phant>'
    ))
    # Real-like b00216/217: Vox_g(S_B) with zeroAsc/zeroDesc only.
    doc.add_paragraph()._p.append(_omath(
        '<m:phant><m:phantPr><m:zeroAsc m:val="1"/><m:zeroDesc m:val="1"/></m:phantPr><m:e>'
        '<m:sSub><m:e><m:r><m:rPr><m:nor/></m:rPr><m:t>Vox</m:t></m:r></m:e>'
        '<m:sub><m:r><m:t>g</m:t></m:r></m:sub></m:sSub>'
        '<m:d><m:e><m:sSub><m:e><m:r><m:t>S</m:t></m:r></m:e>'
        '<m:sub><m:r><m:t>B</m:t></m:r></m:sub></m:sSub></m:e></m:d>'
        '</m:e></m:phant>'
    ))
    hidden = doc.add_paragraph()
    hidden.add_run('keep-visible ')
    hidden._p.append(_omath(
        '<m:phant><m:phantPr><m:show m:val="0"/><m:zeroWid m:val="1"/></m:phantPr>'
        '<m:e><m:r><m:t>AAA</m:t></m:r></m:e></m:phant>'
    ))
    hidden.add_run(' tail')
    off = doc.add_paragraph()
    off.add_run('still ')
    off._p.append(_omath(
        '<m:phant><m:phantPr><m:show m:val="false"/></m:phantPr>'
        '<m:e><m:r><m:t>HIDDEN</m:t></m:r></m:e></m:phant>'
        '<m:phant><m:phantPr><m:show m:val="off"/><m:zeroWid m:val="1"/></m:phantPr>'
        '<m:e><m:r><m:t>OFFHIDDEN</m:t></m:r></m:e></m:phant>'
    ))
    shown = doc.add_paragraph()
    shown._p.append(_omath(
        '<m:phant><m:phantPr><m:show m:val="1"/></m:phantPr>'
        '<m:e><m:r><m:t>KEPT</m:t></m:r></m:e></m:phant>'
        '<m:phant><m:phantPr><m:show m:val="on"/></m:phantPr>'
        '<m:e><m:r><m:t>KEPT2</m:t></m:r></m:e></m:phant>'
        '<m:phant><m:phantPr><m:zeroWid m:val="1"/></m:phantPr>'
        '<m:e><m:r><m:t>WIDTEXT</m:t></m:r></m:e></m:phant>'
    ))
    doc.save(p)

    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    compact = [re.sub(r'\s+', '', t) for t in texts]
    assert any(r'\bar{u}-1' in t or r'\bar{u}–1' in t for t in compact)
    phi = next(t for t in compact if 'Phi' in t or 'Φ' in t)
    assert 'C(' in phi and 'q' in phi
    assert r'\Phi' in phi or 'Φ' in phi
    vox = next(t for t in compact if 'Vox' in t)
    assert 'Vox' in vox and 'g' in vox and 'S' in vox and 'B' in vox
    blob = '\n'.join(texts)
    assert 'AAA' not in blob
    assert 'HIDDEN' not in blob
    assert 'OFFHIDDEN' not in blob
    assert 'keep-visible' in blob and 'tail' in blob
    assert 'still' in blob
    assert 'KEPT' in blob and 'KEPT2' in blob and 'WIDTEXT' in blob
    assert not any(i['id'].startswith('docx-omath') for i in d['issues'])
    assert all(i['resolution'] is None for i in d['issues'])


def test_docx_symbol_font_multiply_keeps_scripts(data_dir):
    """P1: Symbol F0B4 is ×; other known Symbol slots map; unknown font/code is an issue."""
    from docx import Document

    p = data_dir / 'syms.docx'
    doc = Document()
    pack = doc.add_paragraph()
    pack.add_run('each 4')
    pack._p.append(_sym_run('Symbol', 'F0B4'))
    pack.add_run('2 block and 4')
    pack._p.append(_sym_run('Symbol', 'F0B4'))
    pack.add_run('2')
    pack._p.append(_sym_run('Symbol', 'f0b4'))
    pack.add_run('256 texture')
    bits = doc.add_paragraph()
    bits.add_run('128')
    sup = bits.add_run('2')
    sup.font.superscript = True
    bits._p.append(_sym_run('Symbol', 'F0B4'))
    bits.add_run('8 bits and 83')
    sup3 = bits.add_run('3')
    sup3.font.superscript = True
    bits._p.append(_sym_run('Symbol', 'F0B4'))
    bits.add_run('3')
    sup3b = bits.add_run('3')
    sup3b.font.superscript = True
    bare = doc.add_paragraph()
    bare.add_run('n')
    bare._p.append(_sym_run('Symbol', 'B4'))
    bare.add_run('m pm ')
    bare._p.append(_sym_run('Symbol', 'F0B1'))
    bare.add_run(' inf ')
    bare._p.append(_sym_run('Symbol', 'F0A5'))
    unknown = doc.add_paragraph()
    unknown.add_run('before ')
    unknown._p.append(_sym_run('Wingdings', 'F0B4'))
    unknown.add_run(' after ')
    unknown._p.append(_sym_run('Symbol', 'F0E6'))
    unknown.add_run(' end')
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0).paragraphs[0]
    cell.add_run('276')
    cell_sup = cell.add_run('2')
    cell_sup.font.superscript = True
    cell._p.append(_sym_run('Symbol', 'F0B4'))
    cell.add_run('35')
    doc.save(p)

    d = import_document(p)
    texts = [b['text'] for b in d['blocks']]
    pack_t = next(t for t in texts if 'texture' in t)
    assert '4×2' in pack_t
    assert '4×2×256' in pack_t
    assert '42' not in pack_t
    bits_t = next(t for t in texts if 'bits and' in t)
    assert '128²×8' in bits_t
    assert '83³×3³' in bits_t
    assert '1282' not in bits_t and '833' not in bits_t
    other = next(t for t in texts if t.startswith('n'))
    assert 'n×m' in other
    assert '±' in other and '∞' in other
    unk = next(t for t in texts if 'before' in t)
    assert '×' not in unk
    assert '[sym Wingdings F0B4]' in unk
    assert '[sym Symbol F0E6]' in unk
    assert 'before' in unk and 'after' in unk and 'end' in unk
    cell_t = next(t for t in texts if '276' in t)
    assert '276²×35' in cell_t
    sym_issues = [i for i in d['issues'] if i['id'].startswith('docx-sym-')]
    assert len(sym_issues) == 2
    assert all(i['resolution'] is None for i in sym_issues)
    blob = '\n'.join(i['message'] for i in sym_issues)
    assert 'Wingdings' in blob and 'F0B4' in blob
    assert 'F0E6' in blob
    assert not any('F0B1' in i['message'] or 'F0A5' in i['message'] for i in sym_issues)


W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def test_docx_vanish_hidden_run_omitted_from_formula(data_dir):
    """Hidden w:vanish runs (including OMML) are omitted; vanish=0/false/off stays visible."""
    from docx import Document
    from docx.oxml import parse_xml

    p = data_dir / 'vanish.docx'
    doc = Document()
    # Real appendix fragment: trailing hidden "at ma x  : " after visible R}.
    real = doc.add_paragraph()
    real._p.append(_omath(
        '<m:r><m:t>hkp</m:t></m:r>'
        '<m:r><m:t>∈</m:t></m:r>'
        '<m:r><m:t>1,…,R</m:t></m:r>'
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
        f'<w:vanish/></w:rPr><m:t xml:space="preserve">at ma x  : </m:t></m:r>'
    ))
    toggles = doc.add_paragraph()
    toggles._p.append(_omath(
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:vanish w:val="0"/></w:rPr><m:t>VISIBLE0</m:t></m:r>'
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:vanish w:val="false"/></w:rPr><m:t>VISIBLEFALSE</m:t></m:r>'
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:vanish w:val="off"/></w:rPr><m:t>VISIBLEOFF</m:t></m:r>'
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:vanish w:val="1"/></w:rPr><m:t>HIDDEN1</m:t></m:r>'
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:vanish w:val="true"/></w:rPr><m:t>HIDDENTRUE</m:t></m:r>'
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:vanish w:val="on"/></w:rPr><m:t>HIDDENON</m:t></m:r>'
        '<m:r><m:t>KEPT</m:t></m:r>'
        f'<m:r><w:rPr xmlns:w="{W_NS}"><w:vanish/></w:rPr><m:t>HIDDENXYZ</m:t></m:r>'
    ))
    wr = doc.add_paragraph()
    wr.add_run('shown ')
    hidden_run = wr.add_run('secret vanish')
    hidden_run._r.get_or_add_rPr().append(
        parse_xml(f'<w:vanish xmlns:w="{W_NS}"/>')
    )
    wr.add_run(' tail')
    off_run = doc.add_paragraph()
    off_run.add_run('keep-run ')
    visible_run = off_run.add_run('EXPLICITOFF')
    visible_run._r.get_or_add_rPr().append(
        parse_xml(f'<w:vanish xmlns:w="{W_NS}" w:val="0"/>')
    )
    off_run.add_run(' after')
    doc.save(p)

    d = import_document(p)
    blob = '\n'.join(b['text'] for b in d['blocks'])
    assert 'hkp' in blob and '1' in blob and 'R' in blob
    assert 'at ma x' not in blob
    assert 'at ma' not in blob
    assert 'HIDDENXYZ' not in blob
    assert 'HIDDEN1' not in blob
    assert 'HIDDENTRUE' not in blob
    assert 'HIDDENON' not in blob
    assert 'secret vanish' not in blob
    assert 'VISIBLE0' in blob
    assert 'VISIBLEFALSE' in blob
    assert 'VISIBLEOFF' in blob
    assert 'KEPT' in blob
    assert 'shown' in blob and 'tail' in blob
    assert 'keep-run' in blob and 'EXPLICITOFF' in blob and 'after' in blob
    assert all(i['resolution'] is None for i in d['issues'])


def _write_ole_docx(path, preview_name, preview_bytes):
    from docx import Document
    from lxml import etree
    import zipfile
    import io

    doc = Document()
    doc.add_paragraph('Figure 2. Construction of the spatial hash.')
    doc.add_paragraph('Body after the object.')
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    with zipfile.ZipFile(buf, 'r') as z:
        files = {name: z.read(name) for name in z.namelist()}

    R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    V = 'urn:schemas-microsoft-com:vml'
    O = 'urn:schemas-microsoft-com:office:office'
    PKG_R = 'http://schemas.openxmlformats.org/package/2006/relationships'
    CT = 'http://schemas.openxmlformats.org/package/2006/content-types'

    root = etree.fromstring(files['word/document.xml'])
    body = root.find(f'{{{W_NS}}}body')
    sect = body.find(f'{{{W_NS}}}sectPr')
    p_el = etree.Element(f'{{{W_NS}}}p')
    r_el = etree.SubElement(p_el, f'{{{W_NS}}}r')
    obj = etree.SubElement(r_el, f'{{{W_NS}}}object')
    obj.set(f'{{{W_NS}}}dxaOrig', '5564')
    obj.set(f'{{{W_NS}}}dyaOrig', '2849')
    shape = etree.SubElement(obj, f'{{{V}}}shape')
    shape.set('id', '_x0000_i1025')
    shape.set('style', 'width:239.25pt;height:123pt')
    imagedata = etree.SubElement(shape, f'{{{V}}}imagedata')
    imagedata.set(f'{{{R}}}id', 'rId17')
    ole = etree.SubElement(obj, f'{{{O}}}OLEObject')
    ole.set('Type', 'Embed')
    ole.set('ProgID', 'Word.Picture.8')
    ole.set('ShapeID', '_x0000_i1025')
    ole.set(f'{{{R}}}id', 'rId18')
    if sect is not None:
        sect.addprevious(p_el)
    else:
        body.append(p_el)
    files['word/document.xml'] = etree.tostring(
        root, xml_declaration=True, encoding='UTF-8', standalone=True
    )

    rels = etree.fromstring(files['word/_rels/document.xml.rels'])
    def add_rel(rid, typ, target):
        rel = etree.SubElement(rels, f'{{{PKG_R}}}Relationship')
        rel.set('Id', rid)
        rel.set('Type', typ)
        rel.set('Target', target)
    add_rel(
        'rId17',
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image',
        f'media/{preview_name}',
    )
    add_rel(
        'rId18',
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject',
        'embeddings/oleObject1.bin',
    )
    files['word/_rels/document.xml.rels'] = etree.tostring(
        rels, xml_declaration=True, encoding='UTF-8', standalone=True
    )
    files[f'word/media/{preview_name}'] = preview_bytes
    files['word/embeddings/oleObject1.bin'] = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64

    ctypes = etree.fromstring(files['[Content_Types].xml'])
    def add_override(part_name, ctype):
        el = etree.SubElement(ctypes, f'{{{CT}}}Override')
        el.set('PartName', part_name)
        el.set('ContentType', ctype)
    ctype = 'image/x-emf' if preview_name.endswith('.emf') else 'image/png'
    add_override(f'/word/media/{preview_name}', ctype)
    add_override(
        '/word/embeddings/oleObject1.bin',
        'application/vnd.openxmlformats-officedocument.oleObject',
    )
    files['[Content_Types].xml'] = etree.tostring(
        ctypes, xml_declaration=True, encoding='UTF-8', standalone=True
    )

    with zipfile.ZipFile(path, 'w') as z:
        for name, data in files.items():
            z.writestr(name, data)


def test_docx_ole_word_picture_not_silent(data_dir):
    """Word.Picture OLE with VML preview is an issue (and/or extracted media), never omitted."""
    from reader.store import folder

    p = data_dir / 'ole.docx'
    _write_ole_docx(p, 'image7.emf', b'\x01\x00\x00\x00l\x00\x00\x00' + b'\x00' * 64)
    d = import_document(p)
    ole_issues = [i for i in d['issues'] if str(i.get('id', '')).startswith('docx-ole-')]
    figs = [b for b in d['blocks'] if b['kind'] == 'figure']
    assert ole_issues, 'OLE Word.Picture must not be dropped without an issue'
    assert all(i['resolution'] is None for i in ole_issues)
    blob = '\n'.join(i['message'] for i in ole_issues)
    assert 'Word.Picture.8' in blob
    assert 'rId17' in blob and 'rId18' in blob
    assert 'image7.emf' in blob
    assert figs, 'OLE object should occupy a figure slot in reading order'
    directory = folder(d['id'])
    retained = list(directory.glob('image-*.emf')) + [
        directory / b['asset'] for b in figs if b.get('asset')
    ]
    assert any(path.is_file() and path.stat().st_size > 0 for path in retained)
    assert any('Figure 2' in b['text'] for b in d['blocks'])
    assert any('Body after the object' in b['text'] for b in d['blocks'])


def test_docx_srcrect_nonzero_crop_applied_or_issued(data_dir):
    """Non-zero a:srcRect is applied when saving, or recorded as an unresolved limitation."""
    from docx import Document
    from lxml import etree
    from PIL import Image as PILImage
    from reader.store import folder
    import zipfile

    img = data_dir / 'stripes.png'
    im = PILImage.new('RGB', (10, 4), color=(255, 0, 0))
    for x in range(5, 10):
        for y in range(4):
            im.putpixel((x, y), (0, 0, 255))
    im.save(img)

    p = data_dir / 'crop.docx'
    doc = Document()
    doc.add_paragraph().add_run().add_picture(str(img))
    doc.save(p)

    A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    with zipfile.ZipFile(p, 'r') as z:
        files = {name: z.read(name) for name in z.namelist()}
    root = etree.fromstring(files['word/document.xml'])
    for blip in root.iter(f'{{{A}}}blip'):
        rect = etree.Element(f'{{{A}}}srcRect')
        rect.set('l', '50000')
        blip.addnext(rect)
    files['word/document.xml'] = etree.tostring(
        root, xml_declaration=True, encoding='UTF-8', standalone=True
    )
    with zipfile.ZipFile(p, 'w') as z:
        for name, data in files.items():
            z.writestr(name, data)

    d = import_document(p)
    figs = [b for b in d['blocks'] if b['kind'] == 'figure' and b.get('asset')]
    assert figs
    directory = folder(d['id'])
    out = PILImage.open(directory / figs[0]['asset'])
    mid = out.getpixel((out.size[0] // 2, out.size[1] // 2))
    cropped = out.size[0] < 10 and mid[2] > 200 and mid[0] < 80
    issued = any(
        'srcRect' in i['message'] or 'crop' in i['message'].lower()
        for i in d['issues']
    )
    assert cropped or issued
    assert all(i['resolution'] is None for i in d['issues'])

